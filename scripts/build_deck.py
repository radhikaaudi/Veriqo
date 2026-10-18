"""Build the Veriqo AI Builder Cup deck on the official template layout.

    python scripts/build_deck.py            (needs python-pptx)

Every product screenshot is a real capture of the deployed app (demo-assets/screenshots/);
every number is a measured result (docs/05-TEST-RESULTS.md, docs/05b-MODEL-REVIEW.md).
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
A = ROOT / "demo-assets/ppt/template-assets"
SHOT = ROOT / "demo-assets/screenshots"
OUT = ROOT / "docs/Veriqo_AI_Builder_Cup_Deck.pptx"

LIVE_URL = "https://veriqo-api-920194228751.us-central1.run.app"
REPO_URL = "https://github.com/radhikaaudi/Veriqo"
VIDEO_URL = "https://youtu.be/zZbVIwDWArE"

FONT = "Arial"
BLACK, WHITE = RGBColor(0, 0, 0), RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x2A, 0x28, 0x2A)
MUTED = RGBColor(0x5F, 0x62, 0x66)
MAUVE = RGBColor(0x73, 0x4B, 0x5E)
BLUE = RGBColor(0x3F, 0x6E, 0xA3)
GREEN = RGBColor(0x2F, 0x8A, 0x4E)
AMBER = RGBColor(0xA9, 0x70, 0x0C)
RED = RGBColor(0xC0, 0x39, 0x2B)
CHAR = RGBColor(0x4A, 0x4C, 0x4B)
OLIVE = RGBColor(0x74, 0x77, 0x7B)
PANEL = RGBColor(0xF6, 0xF3, 0xF2)
LINE = RGBColor(0xD6, 0xD2, 0xD0)
WARN_BG = RGBColor(0xFD, 0xF3, 0xE4)
RED_BG = RGBColor(0xFC, 0xEC, 0xEA)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, l, t, w, h, fill=None, line=None, line_w=None, shape=MSO_SHAPE.RECTANGLE):
    sp = s.shapes.add_shape(shape, l, t, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = line_w or Pt(1)
    sp.shadow.inherit = False
    return sp


def txt(s, text, l, t, w, h, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, ls=1.08):
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = ls
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = color
    return tb


def bullets(s, items, l, t, w, h, size=15, gap=10):
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.08
        lead, rest = it if isinstance(it, tuple) else ("", it)
        r = p.add_run(); r.text = "•  " + lead
        r.font.bold = bool(lead); r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = INK
        if rest:
            r2 = p.add_run(); r2.text = rest
            r2.font.size = Pt(size); r2.font.name = FONT; r2.font.color.rgb = INK
    return tb


def content(title):
    s = slide()
    s.shapes.add_picture(str(A / "content_header.png"), 0, 0, width=SW)
    rect(s, 0, SH - Inches(0.14), SW, Inches(0.14), fill=BLACK)
    txt(s, title, Inches(0.55), Inches(0.95), Inches(12.2), Inches(0.7), size=27, bold=True)
    return s


def box(s, l, t, w, h, title, body="", fill=PANEL, border=LINE, ts=12, bs=10, color=INK):
    sp = rect(s, l, t, w, h, fill=fill, line=border, line_w=Pt(1.25), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Pt(7); tf.margin_top = tf.margin_bottom = Pt(4)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title; r.font.bold = True; r.font.size = Pt(ts); r.font.name = FONT; r.font.color.rgb = color
    if body:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = body; r2.font.size = Pt(bs); r2.font.name = FONT; r2.font.color.rgb = color
    return sp


def note(s, text, t, fill=WARN_BG, border=AMBER, size=12.5, h=Inches(0.7)):
    sp = rect(s, Inches(0.55), t, Inches(12.2), h, fill=fill, line=border, line_w=Pt(1.25), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Pt(12)
    p = tf.paragraphs[0]; r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = INK
    return sp


def arrow(s, l, t, w, h=Inches(0.28)):
    return rect(s, l, t, w, h, fill=OLIVE, shape=MSO_SHAPE.RIGHT_ARROW)


def down(s, l, t, h, w=Inches(0.3)):
    return rect(s, l, t, w, h, fill=OLIVE, shape=MSO_SHAPE.DOWN_ARROW)


def chip(s, l, t, label, color, w=Inches(2.0)):
    sp = rect(s, l, t, w, Inches(0.32), fill=color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = sp.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; r.font.size = Pt(11); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = WHITE


def column(file):
    """Crop a 1440-px-wide capture to the app's centred content column (no resampling of content)."""
    from PIL import Image
    out = ROOT / "demo-assets/ppt/crops" / file
    out.parent.mkdir(parents=True, exist_ok=True)
    im = Image.open(SHOT / file)
    k = im.width / 1440                      # captures are 1440 CSS px wide at 2x
    im.crop((int(270 * k), 0, int(1170 * k), im.height)).save(out)
    return out


def shot(s, file, l, t, w, caption=None, crop=False):
    """Real screenshot with a thin frame; height follows the image's aspect ratio."""
    src = column(file) if crop else SHOT / file
    pic = s.shapes.add_picture(str(src), l, t, width=w)
    pic.line.color.rgb = LINE; pic.line.width = Pt(1)
    if caption:
        txt(s, caption, l, t + pic.height + Inches(0.08), w, Inches(0.4), size=12, bold=True, align=PP_ALIGN.CENTER)
    return pic


def table(s, rows, l, t, widths, size=12.5, h=Inches(0.42)):
    tb = s.shapes.add_table(len(rows), len(widths), l, t, sum(widths, Inches(0)), h * len(rows)).table
    for i, w in enumerate(widths):
        tb.columns[i].width = w
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tb.cell(ri, ci); c.text = val
            run = c.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(size); run.font.name = FONT
            run.font.bold = ri == 0
            run.font.color.rgb = WHITE if ri == 0 else INK
            c.fill.solid(); c.fill.fore_color.rgb = INK if ri == 0 else (PANEL if ri % 2 == 0 else WHITE)
    return tb


# ====================== 1 · TITLE / TEAM DETAILS ======================
s = slide()
rect(s, 0, 0, SW, SH, fill=WHITE)
s.shapes.add_picture(str(A / "title_band.png"), 0, 0, width=SW)
ty = Inches(4.25)
txt(s, "Veriqo", Inches(0.7), ty, Inches(7), Inches(0.8), size=40, bold=True)
txt(s, "Know what to check before you buy.", Inches(0.72), ty + Inches(0.82), Inches(9), Inches(0.5), size=18, bold=True, color=MAUVE)
tb = s.shapes.add_textbox(Inches(0.72), ty + Inches(1.5), Inches(5.6), Inches(1.6)); tf = tb.text_frame
for i, (k, v) in enumerate([("Team name:  ", "Obscura"), ("Team leader:  ", "Radhika Audichya"), ("Team member:  ", "Somya Kedia")]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(4)
    r = p.add_run(); r.text = k; r.font.bold = True; r.font.size = Pt(15); r.font.name = FONT; r.font.color.rgb = MUTED
    r2 = p.add_run(); r2.text = v; r2.font.size = Pt(15); r2.font.name = FONT; r2.font.color.rgb = INK
txt(s, "Problem Statement", Inches(6.8), ty + Inches(1.45), Inches(6), Inches(0.4), size=14, bold=True, color=MUTED)
txt(s, "Used-car buyers can't quickly tell what a car's photos actually show, what its registration "
       "record says, or what still needs checking — so they overpay or buy hidden problems. Veriqo gives "
       "a fast, evidence-based pre-purchase report that is honest about what it could not verify.",
    Inches(6.8), ty + Inches(1.85), Inches(5.9), Inches(1.5), size=13.5, ls=1.12)

# ====================== 2 · BRIEF ======================
s = content("Brief about the idea")
txt(s, "Veriqo turns a used-car buyer's photos, the registration number and the seller's listing into one "
       "evidence-tagged report — measured median 9.8 seconds from upload to report on the live deployment.",
    Inches(0.55), Inches(1.75), Inches(12.2), Inches(0.9), size=17, ls=1.15)
bullets(s, [
    ("Gemini 2.5 Flash reads every photo: ", "visible damage, rust, tyre wear and interior damage — each finding linked to the photo it came from."),
    ("A registration lookup adds record facts ", "(insurance, loan, challans, blacklist). In this prototype it is SIMULATED with synthetic test plates; any other plate is shown as “not verified”."),
    ("Deterministic rules ", "build a transparent score, flag a listing that doesn't match the record, and flag missing photo angles."),
    ("Every finding is labelled ", "Observed · Record · Inferred · Not checked — and every report ends with a prioritised “what to check before paying” list."),
], Inches(0.55), Inches(2.85), Inches(12.2), Inches(2.8), size=15, gap=10)
note(s, "Veriqo is a red-flag triage. It never certifies a car as safe — accident history and hidden mechanical "
        "condition are always listed as Not checked, so a clean-looking car never gets an “all clear”.",
     Inches(5.9), h=Inches(0.8))

# ====================== 3 · OPPORTUNITIES ======================
s = content("Opportunities")
bullets(s, [
    ("How different is it?  ", "Marketplaces inspect only the cars they sell; independent inspections are paid, physical and take a day. "
     "Veriqo is an instant, self-serve first check for any car — including a private listing."),
    ("How does it solve the problem?  ", "Before any money or travel, the buyer sees what the photos show, what the record says, "
     "what doesn't match, and exactly what to inspect."),
    ("USP — evidence you can trace.  ", "Every finding says where it came from and what it's based on; the score is the sum of "
     "published rules; anything simulated or unverified is labelled as such."),
], Inches(0.55), Inches(1.9), Inches(12.2), Inches(3.2), size=16, gap=16)
txt(s, "The four provenance labels used in every report:", Inches(0.55), Inches(5.3), Inches(12), Inches(0.4), size=13, bold=True, color=MUTED)
for i, (lab, col) in enumerate([("Observed", MAUVE), ("Record", BLUE), ("Inferred", CHAR), ("Not checked", OLIVE)]):
    chip(s, Inches(0.55 + 2.2 * i), Inches(5.75), lab, col)

# ====================== 4 · FEATURES ======================
s = content("List of features offered by the solution")
bullets(s, [
    ("Photo analysis with evidence. ", "Each finding shows the photo it came from, with Gemini's confidence."),
    ("Registration check. ", "Insurance, loan, challans, blacklist — simulated in this prototype and labelled so."),
    ("Listing vs record check. ", "Flags when the seller's make/model differs from the record."),
    ("Missing-angle detection. ", "Says which views (rear, sides, tyres, interior) were never shown."),
], Inches(0.55), Inches(1.9), Inches(6.0), Inches(4.0), size=14, gap=13)
bullets(s, [
    ("Transparent score. ", "Sum of published rules; photos alone can't max it out (capped at 60)."),
    ("Prioritised checklist. ", "“What to check before paying”, highest-impact first."),
    ("Saved, shareable report. ", "Stored in Firestore; the link reopens it after a refresh."),
    ("Safe failure handling. ", "Rejects non-images and >12 MB files; clear messages if the model or network fails."),
], Inches(6.85), Inches(1.9), Inches(6.0), Inches(4.0), size=14, gap=13)
note(s, "Honesty built in: simulated registration data is labelled “SIMULATED — NOT GOVERNMENT VERIFIED”; "
        "an unknown plate is shown as “not verified” — never filled with invented data.", Inches(6.0))

# ====================== 5 · PROCESS FLOW ======================
s = content("Process flow diagram")
steps = [("1. Buyer input", "Photos + registration no. + listing make/model"),
         ("2. Validate & store", "Real-image check · private Cloud Storage"),
         ("3. Gemini vision", "Structured findings per photo (Observed)"),
         ("4. Registration lookup", "Record facts — simulated in this prototype"),
         ("5. Rules engine", "Cross-checks · coverage · score (deterministic)"),
         ("6. Report", "4 labels + checklist → Firestore + audit trail")]
cols = [GREEN, OLIVE, MAUVE, BLUE, CHAR, INK]
gap = Inches(0.22); bw = (Inches(12.2) - gap * 5) / 6; x = Inches(0.55); top = Inches(2.85)
for i, (t, b) in enumerate(steps):
    box(s, x, top, bw, Inches(1.8), t, b, border=cols[i], ts=12, bs=9.5)
    if i < 5:
        arrow(s, x + bw + Inches(0.01), top + Inches(0.76), gap - Inches(0.02))
    x += bw + gap
txt(s, "One buyer action in, one traceable report out. If Gemini fails, nothing is saved and the buyer sees a clear retry message.",
    Inches(0.55), Inches(5.2), Inches(12.2), Inches(0.6), size=14, color=MUTED)

# ====================== 6 · WIREFRAMES → REAL SCREENS ======================
s = content("Wireframes / Mock diagrams of the proposed solution")
shot(s, "02-upload-filled.png", Inches(0.55), Inches(1.85), Inches(5.95), "Upload: plate, listing, numbered photos")
shot(s, "05-findings.png", Inches(6.8), Inches(1.85), Inches(5.95), "Report: findings labelled by source, with photo evidence")
txt(s, f"Actual screens captured from the deployed app ({LIVE_URL.replace('https://', '')}) — not mockups.",
    Inches(0.55), Inches(6.55), Inches(12.2), Inches(0.4), size=12, color=MUTED)

# ====================== 7 · ARCHITECTURE ======================
s = content("Architecture diagram of the proposed solution")
cx, cw = Inches(3.65), Inches(6.0)
box(s, cx, Inches(1.75), cw, Inches(0.7), "Buyer's browser", "React web app", ts=13)
down(s, cx + cw / 2 - Inches(0.15), Inches(2.5), Inches(0.35))
box(s, cx, Inches(2.9), cw, Inches(0.85), "Cloud Run · veriqo-api (us-central1)",
    "Serves the web app + FastAPI · image validation · deterministic rules", fill=RGBColor(0xEE, 0xF5, 0xF0), border=GREEN, ts=13)
down(s, cx + cw / 2 - Inches(0.15), Inches(3.8), Inches(0.45))
items = [("Gemini 2.5 Flash", "Vertex AI · structured findings", MAUVE, PANEL),
         ("Registration lookup", "SIMULATED — synthetic test fixtures", AMBER, WARN_BG),
         ("Cloud Firestore", "Reports + append-only audit trail", GREEN, PANEL),
         ("Cloud Storage", "Private bucket · 30-day expiry", BLUE, PANEL)]
x = Inches(0.55)
for t, b, c, f in items:
    box(s, x, Inches(4.35), Inches(2.9), Inches(1.1), t, b, fill=f, border=c, ts=13, bs=10)
    x += Inches(3.1)
txt(s, "Gemini perceives; Cloud Run decides with fixed rules; Firestore keeps the auditable record. "
       "Not used in this prototype: Firebase Hosting and Firebase Authentication (Cloud Run serves the app; there is no sign-in).",
    Inches(0.55), Inches(5.75), Inches(12.2), Inches(0.9), size=13, color=MUTED, ls=1.12)

# ====================== 8 · TECHNOLOGIES ======================
s = content("Technologies to be used in the solution")
bullets(s, [
    ("Gemini 2.5 Flash on Vertex AI — ", "multimodal photo analysis with a typed JSON response schema (temperature 0)."),
    ("Cloud Run — ", "one service serving the React app and the FastAPI backend; 1 warm instance."),
    ("Cloud Firestore (Firebase) — ", "saved reports and an append-only audit trail."),
    ("Cloud Storage — ", "uploaded photos in a private bucket, deleted automatically after 30 days."),
    ("Cloud Build + Artifact Registry — ", "builds and stores the container on every deploy."),
    ("React + Vite + TypeScript · FastAPI + Pydantic (Python) — ", "app and API."),
    ("pytest + Playwright — ", "API tests and browser tests run against the live deployment."),
    ("Registration data — ", "synthetic fixtures today; built to plug in a licensed provider (Surepass / Signzy class)."),
], Inches(0.55), Inches(1.85), Inches(12.2), Inches(5.0), size=14.5, gap=9)

# ====================== 9 · COST ======================
s = content("Estimated implementation cost (optional)")
txt(s, "Indicative, prototype scale. These are estimates, not measured bills.", Inches(0.55), Inches(1.8), Inches(12.2), Inches(0.5), size=14, color=MUTED)
table(s, [("Component", "Driver", "Indicative cost"),
          ("Gemini 2.5 Flash", "≈5 photos per report", "Fractions of ₹ per report"),
          ("Cloud Run", "1 always-warm instance (avoids cold starts)", "Small fixed monthly cost"),
          ("Firestore + Cloud Storage", "Prototype volumes; photos auto-deleted at 30 days", "Within free tier"),
          ("Registration provider (when live)", "Per-lookup fee", "Likely the largest per-report cost"),
          ],
      Inches(0.55), Inches(2.5), [Inches(3.6), Inches(5.3), Inches(3.3)], size=13)

# ====================== 10 · SNAPSHOTS ======================
s = content("Snapshots of the prototype")
shot(s, "04-report-summary.png", Inches(0.55), Inches(1.7), Inches(3.95), "Score, breakdown, simulated label", crop=True)
shot(s, "06-checklist.png", Inches(4.69), Inches(1.7), Inches(3.95), "Prioritised buyer checklist", crop=True)
shot(s, "07-rc-not-verified.png", Inches(8.83), Inches(1.7), Inches(3.95), "Unknown plate → not verified", crop=True)
txt(s, "Captured from the live deployment with real Gemini output. Test photos: openly licensed Wikimedia Commons images "
       "(credits in the repo). Registration data is synthetic and labelled as simulated.",
    Inches(0.55), Inches(6.55), Inches(12.2), Inches(0.5), size=11.5, color=MUTED)

# ====================== 11 · BENCHMARKING ======================
s = content("Prototype Performance report/Benchmarking")
table(s, [("What we measured", "Result"),
          ("API tests (pytest, 21 cases incl. invalid files, model failures, persistence)", "21 / 21 pass"),
          ("Scoring-rule evaluation (11 labelled cases)", "67 / 67 checks pass"),
          ("Browser tests on the live deployment (15 scenarios)", "15 / 15 pass  (first run: 14 / 15)"),
          ("Upload → report time, 13 live runs", "median 9.8 s · range 6.8–17.2 s")],
      Inches(0.55), Inches(1.8), [Inches(8.4), Inches(3.8)], size=13)
txt(s, "Model review (checked by hand, before → after prompt and rule fixes)", Inches(0.55), Inches(4.15), Inches(12.2), Inches(0.4), size=14, bold=True)
bullets(s, [
    ("Fixed: ", "one key scratch scored 100/100 → now 26 (one moderate item); sharp close-ups were marked low-confidence → now 0.7–1.0; "
                "a blurry photo claimed views → now claims none; repeat runs disagreed → now identical."),
    ("Still open: ", "headlight severity varies between photo sets; a dent photo still yields 6 items. "
                     "No precision/recall is claimed — 11 photos are not a labelled benchmark."),
], Inches(0.55), Inches(4.6), Inches(12.2), Inches(1.6), size=13, gap=7)
txt(s, "Sources: docs/05-TEST-RESULTS.md · docs/05b-MODEL-REVIEW.md", Inches(0.55), Inches(6.55), Inches(12.2), Inches(0.4), size=11.5, color=MUTED)

# ====================== 12 · FUTURE ======================
s = content("Additional Details/Future Development")
bullets(s, [
    ("Live registration data. ", "Connect a licensed provider so records become “Record verified”, not simulated."),
    ("A real accuracy benchmark. ", "~100 labelled photo sets checked by mechanics → per-defect precision/recall."),
    ("Accounts. ", "Firebase Authentication so buyers can keep and share their reports."),
    ("History signals. ", "Partner data for accident, service and odometer history."),
    ("Local languages. ", "Reports in Hindi and regional languages with Gemini."),
    ("Same engine, new listings. ", "Verify rental-flat listings with the same evidence-tagged approach."),
], Inches(0.55), Inches(1.9), Inches(12.2), Inches(4.6), size=15, gap=11)

# ====================== 13 · LINKS ======================
s = content("Provide links to your")
yy = Inches(2.1)
for lbl, val in [("GitHub Public Repository", REPO_URL),
                 ("Demo Video Link (3 Minutes)", VIDEO_URL),
                 ("Final Product Link", LIVE_URL)]:
    txt(s, "●  " + lbl, Inches(0.8), yy, Inches(4.9), Inches(0.5), size=17, bold=True)
    rect(s, Inches(5.7), yy - Inches(0.05), Inches(7.05), Inches(0.55), fill=PANEL, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    link = txt(s, val, Inches(5.9), yy + Inches(0.07), Inches(6.7), Inches(0.4), size=13, color=BLUE)
    link.text_frame.paragraphs[0].runs[0].hyperlink.address = val   # clickable in PowerPoint and the PDF
    yy += Inches(1.0)
note(s, "The live app is kept warm (1 always-on Cloud Run instance) for the evaluation period.", Inches(5.4), fill=PANEL, border=LINE)

# ====================== 14 · THANK YOU ======================
s = slide()
rect(s, 0, 0, SW, SH, fill=BLACK)
s.shapes.add_picture(str(A / "thankyou.png"), 0, 0, width=SW, height=SH)

prs.save(OUT)
print("saved", OUT, "·", len(prs.slides._sldIdLst), "slides")
